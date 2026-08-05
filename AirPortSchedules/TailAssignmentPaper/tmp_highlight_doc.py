from docx import Document
from docx.shared import RGBColor
from shutil import copy2

src = 'Integrated aircraft routing & maintenance Planning_v0.011.docx'
bak = src + '.bak'
copy2(src, bak)
doc = Document(src)

targets = [
    'deadhead', 'dead-head', 'maintenance station', 'resource capacity',
    'this way we will provide', 'the model will then optimize',
    'the above model can be re-run', 'thus, our solution',
    'the exact problem', 'the inputs for th model', 'the airline industry',
]

for p in doc.paragraphs:
    txt = p.text.lower()
    if any(k in txt for k in targets):
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            run.font.bold = True
        p.alignment = 1

# also mark any paragraph mentioning 'aircraft routing' broadly if it seems generic
for p in doc.paragraphs:
    txt = p.text.lower()
    if 'aircraft routing' in txt and 'maintenance' in txt and len(txt.split()) > 8:
        for run in p.runs:
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            run.font.bold = True
        p.alignment = 1

doc.save(src)
print('updated', src)
print('backup', bak)
